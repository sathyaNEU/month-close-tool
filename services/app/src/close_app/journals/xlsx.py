"""XLSX import (Stage 7b): a workbook's first sheet becomes the same rows a CSV import reads.

A workbook is a zip file of XML, and ``openpyxl`` alone doesn't guard against XML entity expansion (``defusedxml``
isn't installed), so nothing parses one before these checks pass:

- **The zip** has at most 200 entries and 50 MiB uncompressed in all, and no entry is compressed more than 100 to
  one, which stops zip bombs.
- **Parsing runs in a subprocess** with a CPU limit, a timeout and, on Linux, a memory limit, so an XML bomb can
  only take that process down. One workbook is read at a time per app process, and never inside a database
  transaction, so slow or hostile uploads can't use up the app's memory or its connection pool.
- **Cells.** Only cached values are read. A formula with no cached value (a workbook never calculated) is
  refused. Numbers become exact decimals, so a stored ``0.30000000000000004`` stays visible and the import refuses
  its extra decimals. Dates become ``YYYY-MM-DD``.
"""

from __future__ import annotations

import contextlib
import csv
import io
import json
import subprocess
import sys
import threading
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Final

from close_app.errors import Problem

MAX_ENTRIES: Final = 200
MAX_UNCOMPRESSED: Final = 50 * 1024 * 1024
MAX_RATIO: Final = 100
TIMEOUT_S: Final = 20
MEMORY_LIMIT: Final = 256 * 1024 * 1024  # the app container has 512 MiB in all
CPU_SECONDS: Final = 10
SLOT_WAIT_S = 10.0  # how long a second upload waits for the one parsing slot
_SLOTS: Final = threading.BoundedSemaphore(1)
CONTENT_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _refuse(message: str, code: str = "invalid_workbook") -> Problem:
    issue = {"code": code, "path": "/", "message": message}
    return Problem(422, "The workbook was not imported", kind="validation-failed", issues=[issue])


def check_zip(data: bytes) -> None:
    """Refuse anything that isn't a reasonable zip before any XML is parsed."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise _refuse("the file is not an XLSX workbook") from None
    entries = archive.infolist()
    if len(entries) > MAX_ENTRIES:
        raise _refuse(f"the workbook has more than {MAX_ENTRIES} parts", "zip_bomb")
    if sum(e.file_size for e in entries) > MAX_UNCOMPRESSED:
        raise _refuse(f"the workbook expands to more than {MAX_UNCOMPRESSED} bytes", "zip_bomb")
    for e in entries:
        if e.compress_size and e.file_size / e.compress_size > MAX_RATIO:
            raise _refuse(f"{e.filename} is compressed more than {MAX_RATIO} to one", "zip_bomb")


def _limits() -> None:  # pragma: no cover - runs in the child process
    import resource

    resource.setrlimit(resource.RLIMIT_CPU, (CPU_SECONDS, CPU_SECONDS))
    with contextlib.suppress(ValueError, OSError):  # macOS doesn't enforce RLIMIT_AS; the other limits still apply
        resource.setrlimit(resource.RLIMIT_AS, (MEMORY_LIMIT, MEMORY_LIMIT))


def to_csv(data: bytes) -> bytes:
    """The first sheet as CSV bytes, parsed in a limited subprocess; a 422 with the reason otherwise, or a 503 when
    another workbook holds the parsing slot for too long."""
    check_zip(data)
    if not _SLOTS.acquire(timeout=SLOT_WAIT_S):
        raise Problem(503, "Another workbook is being read", "retry shortly")
    try:
        done = subprocess.run(
            [sys.executable, "-m", "close_app.journals.xlsx"],
            input=data,
            capture_output=True,
            timeout=TIMEOUT_S,
            preexec_fn=_limits,
            check=False,
        )
    except subprocess.TimeoutExpired:
        raise _refuse("the workbook took too long to read") from None
    finally:
        _SLOTS.release()
    if done.returncode != 0:
        try:
            reason = str(json.loads(done.stderr.decode("utf-8").strip().splitlines()[-1])["error"])
        except (ValueError, IndexError, KeyError):
            reason = "the workbook could not be read"
        raise _refuse(reason)
    return done.stdout


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(Decimal(str(value)))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _convert(data: bytes) -> bytes:  # pragma: no cover - runs in the child process; tested through to_csv
    from openpyxl import load_workbook

    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet_f, sheet_v = formulas.worksheets[0], values.worksheets[0]
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    for r, (row_f, row_v) in enumerate(zip(sheet_f.iter_rows(), sheet_v.iter_rows(), strict=False), start=1):
        cells = []
        for cell_f, cell_v in zip(row_f, row_v, strict=False):
            if isinstance(cell_f.value, str) and cell_f.value.startswith("=") and cell_v.value is None:
                raise ValueError(f"row {r}: a formula has no calculated value; open and save the workbook in Excel")
            cells.append(_text(cell_v.value))
        while cells and cells[-1] == "":
            cells.pop()
        if cells:
            writer.writerow(cells)
    return out.getvalue().encode("utf-8")


if __name__ == "__main__":  # pragma: no cover - the subprocess entry point
    try:
        sys.stdout.buffer.write(_convert(sys.stdin.buffer.read()))
    except Exception as exc:
        sys.stderr.write(json.dumps({"error": f"{type(exc).__name__}: {exc}"[:500]}) + "\n")
        raise SystemExit(3) from None
